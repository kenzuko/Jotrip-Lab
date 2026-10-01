"""Respectful one-shot public-source probe for Weather Engine V3 research.

The probe does not bypass authentication, does not crawl recursively, does not
archive radar imagery, and does not call discovered private/unknown endpoints.

It only:
1) inspects public iWeather HTML/static JS for route hints;
2) requests the already-public 60018 XLSX export once and summarizes whether
   it contains usable numeric cells.

Output is metadata only.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import sys
import time
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.parse import urlencode, urljoin, urlparse
from urllib.request import Request, urlopen

try:
    import openpyxl
except Exception:
    openpyxl = None

UA = "JoTrip-WeatherV3-Research/1.0 (+public-source-probe; low-rate)"
IWEATHER_URL = "https://iweather.gov.vn/dashboard/?areaRadar=COM&lightning=true&productRadar=CMAX"
KTT_EXPORT_URL = "https://kttvtudong.net/kttv/export/excelexport?sid=33"
OUTPUT = Path("weather/research/v3-public-source-probe.json")

MAX_JS_FILES = 12
MAX_JS_BYTES_TOTAL = 6_000_000
KEYWORDS = ("radar", "cmax", "lightning", "tile", "wms", "qpe", "rain", "api", "geojson", "image")


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def fetch(url: str, *, method: str = "GET", data: bytes | None = None, accept: str = "*/*") -> tuple[bytes, dict[str, Any]]:
    req = Request(
        url,
        data=data,
        method=method,
        headers={
            "User-Agent": UA,
            "Accept": accept,
            "Accept-Language": "vi-VN,vi;q=0.9,en;q=0.5",
        },
    )
    started = time.time()
    with urlopen(req, timeout=30) as res:
        body = res.read()
        meta = {
            "requested_url": url,
            "final_url": res.geturl(),
            "status": getattr(res, "status", None),
            "content_type": res.headers.get("Content-Type"),
            "content_disposition": res.headers.get("Content-Disposition"),
            "bytes": len(body),
            "sha256": sha256(body),
            "elapsed_ms": round((time.time() - started) * 1000),
        }
    return body, meta


class AssetParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.scripts: list[str] = []
        self.iframes: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        amap = {k: v for k, v in attrs}
        if tag == "script" and amap.get("src"):
            self.scripts.append(str(amap["src"]))
        if tag == "iframe" and amap.get("src"):
            self.iframes.append(str(amap["src"]))


def route_hints(text: str) -> list[str]:
    hits: set[str] = set()

    # Absolute URLs.
    for value in re.findall(r"https?://[^\s\"'<>]+", text, flags=re.I):
        low = value.lower()
        if any(k in low for k in KEYWORDS):
            hits.add(value[:500])

    # Quoted relative route-like strings.
    for value in re.findall(r"[\"']([^\"']{3,500})[\"']", text):
        low = value.lower()
        if not any(k in low for k in KEYWORDS):
            continue
        if value.startswith("/") or "://" in value or "{" in value or "tile" in low or "wms" in low:
            hits.add(value)

    return sorted(hits)[:200]


def probe_iweather() -> dict[str, Any]:
    result: dict[str, Any] = {"url": IWEATHER_URL}
    try:
        body, meta = fetch(IWEATHER_URL, accept="text/html,*/*")
        result["document"] = meta
        text = body.decode("utf-8", errors="replace")
        parser = AssetParser()
        parser.feed(text)

        base = meta["final_url"]
        iframe_urls = [urljoin(base, x) for x in parser.iframes]
        script_urls = [urljoin(base, x) for x in parser.scripts]
        result["iframe_urls"] = iframe_urls
        result["script_count"] = len(script_urls)
        result["html_route_hints"] = route_hints(text)

        total = 0
        assets = []
        same_host = urlparse(base).netloc
        for url in script_urls:
            if len(assets) >= MAX_JS_FILES or total >= MAX_JS_BYTES_TOTAL:
                break
            if urlparse(url).netloc not in ("", same_host):
                continue
            try:
                js, js_meta = fetch(url, accept="application/javascript,text/javascript,*/*")
            except Exception as exc:
                assets.append({"url": url, "error": repr(exc)})
                continue
            total += len(js)
            js_text = js.decode("utf-8", errors="replace")
            hints = route_hints(js_text)
            assets.append({
                "url": url,
                "meta": js_meta,
                "route_hints": hints,
            })
            # Low-rate behavior.
            time.sleep(0.25)
        result["static_assets_inspected"] = assets
        result["static_asset_bytes_total"] = total
        result["probe_result"] = "STATIC_DISCOVERY_COMPLETE"
    except Exception as exc:
        result["probe_result"] = "UNAVAILABLE"
        result["error"] = repr(exc)
    return result


def probe_60018() -> dict[str, Any]:
    result: dict[str, Any] = {
        "url": KTT_EXPORT_URL,
        "policy": "ONE_EXPORT_ONLY_DO_NOT_BLOCK_V3",
    }
    today = datetime.now(timezone.utc).astimezone().strftime("%d/%m/%Y")
    data = urlencode({"fd": today, "td": today}).encode("utf-8")
    try:
        body, meta = fetch(
            KTT_EXPORT_URL,
            method="POST",
            data=data,
            accept="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*",
        )
        result["export"] = meta
        result["requested_date"] = today
        if openpyxl is None:
            result["parse_status"] = "OPENPYXL_NOT_AVAILABLE"
            return result

        wb = openpyxl.load_workbook(io.BytesIO(body), data_only=True, read_only=True)
        sheets: list[dict[str, Any]] = []
        numeric_measurement_candidates = 0
        for ws in wb.worksheets:
            nonempty: list[tuple[int, int, Any]] = []
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None and str(cell.value).strip() != "":
                        nonempty.append((cell.row, cell.column, cell.value))
            # Headers/times may be numeric/date-like. Measurement candidates are
            # numeric cells outside the first three rows and after the first two columns.
            numeric = [
                (r, c, v)
                for (r, c, v) in nonempty
                if r > 3 and c > 2 and isinstance(v, (int, float)) and not isinstance(v, bool)
            ]
            numeric_measurement_candidates += len(numeric)
            sheets.append({
                "title": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "nonempty_cell_count": len(nonempty),
                "numeric_measurement_candidate_count": len(numeric),
                "sample_nonempty": [
                    {"row": r, "column": c, "value": str(v)[:120]}
                    for r, c, v in nonempty[:30]
                ],
                "sample_numeric_measurements": [
                    {"row": r, "column": c, "value": v}
                    for r, c, v in numeric[:20]
                ],
            })
        result["sheets"] = sheets
        result["numeric_measurement_candidate_count"] = numeric_measurement_candidates
        result["parse_status"] = (
            "NUMERIC_ROWS_PRESENT" if numeric_measurement_candidates > 0
            else "NO_USABLE_NUMERIC_MEASUREMENTS_FOUND"
        )
    except Exception as exc:
        result["parse_status"] = "UNAVAILABLE"
        result["error"] = repr(exc)
    return result


def main() -> int:
    payload = {
        "schema_version": "weather-v3-public-source-probe-v1",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "mode": "RESEARCH_METADATA_ONLY",
        "rules": {
            "no_auth_bypass": True,
            "no_radar_image_archive": True,
            "no_discovered_endpoint_auto_fetch": True,
            "low_rate": True,
        },
        "iweather": probe_iweather(),
        "kttv_60018": probe_60018(),
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "iweather": payload["iweather"].get("probe_result"),
        "iweather_scripts": payload["iweather"].get("script_count"),
        "kttv_60018": payload["kttv_60018"].get("parse_status"),
        "kttv_numeric_candidates": payload["kttv_60018"].get("numeric_measurement_candidate_count"),
        "output": str(OUTPUT),
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
